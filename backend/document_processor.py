from dataclasses import dataclass
from io import BytesIO
from typing import List
import csv
import io

from pypdf import PdfReader
from docx import Document as DocxDocument


@dataclass
class Chunk:
    text: str
    doc_name: str
    chunk_id: int
    page_number: int  

def extract_text_from_pdf(file_bytes: bytes) -> List[tuple]:
    reader = PdfReader(BytesIO(file_bytes))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append((i, text))
    return pages


def extract_text_from_txt(file_bytes: bytes) -> List[tuple]:
    text = file_bytes.decode("utf-8", errors="ignore")
    return [(-1, text)]


def extract_text_from_docx(file_bytes: bytes) -> List[tuple]:
    doc = DocxDocument(BytesIO(file_bytes))

    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    full_text = "\n".join(parts)
    return [(-1, full_text)] if full_text.strip() else []


def extract_text_from_csv(file_bytes: bytes) -> List[tuple]:
    text = file_bytes.decode("utf-8", errors="ignore")
    reader = csv.DictReader(io.StringIO(text))

    lines = []
    for row in reader:
        line = ", ".join(f"{k.strip()}: {v.strip()}" for k, v in row.items() if k and v)
        if line:
            lines.append(line)

    full_text = "\n".join(lines)
    return [(-1, full_text)] if full_text.strip() else []


def load_document(filename: str, file_bytes: bytes) -> List[tuple]:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif lower.endswith(".txt"):
        return extract_text_from_txt(file_bytes)
    elif lower.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif lower.endswith(".csv"):
        return extract_text_from_csv(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {filename}")


def chunk_text(
    pages: List[tuple],
    doc_name: str,
    chunk_size: int = 800,
    overlap: int = 150,
) -> List[Chunk]:
    chunks: List[Chunk] = []
    chunk_id = 0

    for page_number, page_text in pages:
        page_text = page_text.replace("\n", " ").strip()
        start = 0
        while start < len(page_text):
            end = start + chunk_size
            piece = page_text[start:end].strip()
            if piece:
                chunks.append(
                    Chunk(
                        text=piece,
                        doc_name=doc_name,
                        chunk_id=chunk_id,
                        page_number=page_number,
                    )
                )
                chunk_id += 1
            start += chunk_size - overlap  

    return chunks

def process_document(filename: str, file_bytes: bytes) -> List[Chunk]:
    """Full pipeline: bytes in -> list of Chunk objects out."""
    pages = load_document(filename, file_bytes)
    return chunk_text(pages, doc_name=filename)
